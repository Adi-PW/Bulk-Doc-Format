"""
Renders a parsed ``Question`` into the exact QBG ingestion layout, using the
team's own sample doc (assets/question_template.docx) as the template so that
fonts, borders, row heights and colours are identical to the reference file.
"""
from __future__ import annotations

import copy
import io
import os
from typing import List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu

try:  # package layout (qbg/...) or flat layout (all files next to app.py)
    from .docx_reader import Line, Table
    from .parser import Question, clean_line
except ImportError:
    from docx_reader import Line, Table
    from parser import Question, clean_line


def find_asset(name: str) -> str:
    """Locate an asset whether files live in assets/ or flat beside the code."""
    here = os.path.dirname(os.path.abspath(__file__))
    for cand in (os.path.join(here, "..", "assets", name), os.path.join(here, "assets", name),
                 os.path.join(here, name), os.path.join(here, "..", name)):
        if os.path.exists(cand):
            return os.path.normpath(cand)
    return os.path.join(here, "..", "assets", name)


TEMPLATE_PATH = find_asset("question_template.docx")
_TEMPLATE_BYTES: Optional[bytes] = None

# Value-cell width in the template (8585 dxa) minus cell margins.
_CELL_WIDTH_DXA = 8400
_MAX_IMG_EMU = int(_CELL_WIDTH_DXA / 1440 * 914400 * 0.95)


def _template() -> bytes:
    global _TEMPLATE_BYTES
    if _TEMPLATE_BYTES is None:
        with open(TEMPLATE_PATH, "rb") as fh:
            _TEMPLATE_BYTES = fh.read()
    return _TEMPLATE_BYTES


def _cell_text(tc) -> str:
    return "".join(t.text or "" for t in tc.iter(qn("w:t"))).strip()


def _base_rpr(tc):
    """First run's rPr inside a template cell (font family / size of that field)."""
    r = next(tc.iter(qn("w:r")), None)
    if r is None:
        return None
    rpr = r.find(qn("w:rPr"))
    return copy.deepcopy(rpr) if rpr is not None else None


def _base_ppr(tc):
    p = tc.find(qn("w:p"))
    if p is None:
        return None
    ppr = p.find(qn("w:pPr"))
    return copy.deepcopy(ppr) if ppr is not None else None


def _clear_cell(tc):
    for child in list(tc):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            tc.remove(child)


def _set_flag(rpr, tag: str, on: bool):
    el = rpr.find(qn(tag))
    if on:
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        if tag == "w:u":
            el.set(qn("w:val"), "single")
    elif el is not None:
        rpr.remove(el)


def _order_rpr(rpr):
    """rPr children must follow the schema order or Word/Drive may reject them."""
    order = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
             "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
             "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
             "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
             "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
             "w:eastAsianLayout", "w:specVanish", "w:oMath"]
    rank = {qn(t): i for i, t in enumerate(order)}
    kids = sorted(list(rpr), key=lambda e: rank.get(e.tag, 999))
    for k in list(rpr):
        rpr.remove(k)
    for k in kids:
        rpr.append(k)


class QuestionDocBuilder:
    def __init__(self, category: str, language: str = "English", qtype: str = "SCQ"):
        self.category = category
        self.language = language
        self.qtype = qtype

    # ------------------------------------------------------------------ #
    def build(self, q: Question) -> bytes:
        doc = Document(io.BytesIO(_template()))
        self.doc = doc
        tbl = doc.tables[0]._tbl
        rows = tbl.findall(qn("w:tr"))

        by_label = {}
        option_rows = []
        for tr in rows:
            tcs = tr.findall(qn("w:tc"))
            label = _cell_text(tcs[0]) if tcs else ""
            if label == "Option":
                option_rows.append(tr)
            else:
                by_label.setdefault(label, tr)

        self._fill_plain(by_label.get("Exam Category"), self.category)
        self._fill_plain(by_label.get("Input Language"), self.language)
        self._fill_plain(by_label.get("Type"), self.qtype)
        self._fill_rich(by_label.get("Question"), q.body)
        self._fill_rich(by_label.get("Solution"), q.solution)

        # options: clone the first template option row once per option
        proto = copy.deepcopy(option_rows[0])
        anchor = option_rows[0].getprevious()
        for tr in option_rows:
            tbl.remove(tr)
        for k, opt in enumerate(q.options):
            tr = copy.deepcopy(proto)
            tcs = tr.findall(qn("w:tc"))
            self._fill_plain_tc(tcs[1], "Correct" if k == q.correct_index else "Incorrect")
            self._fill_rich_tc(tcs[2], opt)
            anchor.addnext(tr)
            anchor = tr

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    # ------------------------------------------------------------------ #
    def _value_tc(self, tr):
        tcs = tr.findall(qn("w:tc"))
        return tcs[-1]

    def _fill_plain(self, tr, text: str):
        if tr is not None:
            self._fill_plain_tc(self._value_tc(tr), text)

    def _fill_plain_tc(self, tc, text: str):
        rpr, ppr = _base_rpr(tc), _base_ppr(tc)
        _clear_cell(tc)
        p = OxmlElement("w:p")
        if ppr is not None:
            p.append(ppr)
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        p.append(r)
        tc.append(p)

    def _fill_rich(self, tr, items: List[object]):
        if tr is not None:
            self._fill_rich_tc(self._value_tc(tr), items)

    def _fill_rich_tc(self, tc, items: List[object]):
        rpr, ppr = _base_rpr(tc), _base_ppr(tc)
        _clear_cell(tc)
        wrote_any = False
        for it in items:
            if isinstance(it, Table):
                tc.append(self._nested_table(it, rpr, ppr))
                wrote_any = True
            else:
                tc.append(self._paragraph(it, rpr, ppr))
                wrote_any = True
        # a cell must end with a paragraph (also required after a nested table)
        last = tc[-1] if len(tc) else None
        if not wrote_any or last is None or last.tag != qn("w:p"):
            p = OxmlElement("w:p")
            if ppr is not None:
                p.append(copy.deepcopy(ppr))
            tc.append(p)

    def _paragraph(self, line: Line, base_rpr, base_ppr):
        p = OxmlElement("w:p")
        if base_ppr is not None:
            p.append(copy.deepcopy(base_ppr))
        for run in line.runs:
            if run.image is not None:
                p.append(self._image_run(run))
                continue
            if not run.text:
                continue
            # split on tabs so they become real <w:tab/> elements
            r = OxmlElement("w:r")
            rpr = copy.deepcopy(base_rpr) if base_rpr is not None else OxmlElement("w:rPr")
            el = rpr.find(qn("w:highlight"))
            if el is not None:
                rpr.remove(el)
            _set_flag(rpr, "w:b", run.bold)
            _set_flag(rpr, "w:bCs", run.bold)
            _set_flag(rpr, "w:i", run.italic)
            _set_flag(rpr, "w:iCs", run.italic)
            _set_flag(rpr, "w:u", run.underline)
            if run.superscript or run.subscript:
                va = OxmlElement("w:vertAlign")
                va.set(qn("w:val"), "superscript" if run.superscript else "subscript")
                rpr.append(va)
            _order_rpr(rpr)
            r.append(rpr)
            parts = run.text.split("\t")
            for k, part in enumerate(parts):
                if k:
                    r.append(OxmlElement("w:tab"))
                if part:
                    t = OxmlElement("w:t")
                    t.text = part
                    t.set(qn("xml:space"), "preserve")
                    r.append(t)
            p.append(r)
        return p

    def _image_run(self, run):
        # Let python-docx do the relationship + drawing XML, then lift the run.
        scratch = self.doc.add_paragraph()
        cx, cy = run.image_cx, run.image_cy
        width = None
        if cx and cy:
            if cx > _MAX_IMG_EMU:
                cy = int(cy * _MAX_IMG_EMU / cx)
                cx = _MAX_IMG_EMU
            width = Emu(cx)
        try:
            r = scratch.add_run()
            r.add_picture(io.BytesIO(run.image), width=width)
            r_el = r._r
        except Exception:
            r_el = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = "[image could not be embedded]"
            r_el.append(t)
        scratch._p.getparent().remove(scratch._p)
        return r_el

    def _nested_table(self, table: Table, base_rpr, base_ppr):
        n_cols = max(table.n_cols, 1)
        col_w = _CELL_WIDTH_DXA // n_cols
        tbl = OxmlElement("w:tbl")
        tblPr = OxmlElement("w:tblPr")
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), str(col_w * n_cols))
        tblW.set(qn("w:type"), "dxa")
        tblPr.append(tblW)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        tbl.append(tblPr)
        grid = OxmlElement("w:tblGrid")
        for _ in range(n_cols):
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(col_w))
            grid.append(gc)
        tbl.append(grid)
        for row in table.rows:
            tr = OxmlElement("w:tr")
            used = 0
            for cell in row:
                tc = OxmlElement("w:tc")
                tcPr = OxmlElement("w:tcPr")
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(col_w * cell.grid_span))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)
                if cell.grid_span > 1:
                    gs = OxmlElement("w:gridSpan")
                    gs.set(qn("w:val"), str(cell.grid_span))
                    tcPr.append(gs)
                if cell.v_merge:
                    vm = OxmlElement("w:vMerge")
                    if cell.v_merge == "restart":
                        vm.set(qn("w:val"), "restart")
                    tcPr.append(vm)
                tc.append(tcPr)
                lines = [l for l in cell.lines if not l.is_blank()]
                for ln in lines:
                    tc.append(self._paragraph(clean_line(copy.deepcopy(ln)), base_rpr, base_ppr))
                if not lines:
                    tc.append(OxmlElement("w:p"))
                tr.append(tc)
                used += cell.grid_span
            # pad ragged rows so the grid stays rectangular
            while used < n_cols:
                tc = OxmlElement("w:tc")
                tcPr = OxmlElement("w:tcPr")
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(col_w))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)
                tc.append(tcPr)
                tc.append(OxmlElement("w:p"))
                tr.append(tc)
                used += 1
            tbl.append(tr)
        return tbl
