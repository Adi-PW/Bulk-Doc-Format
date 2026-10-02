"""
Robustness suite: builds synthetic .docx files covering layouts that are NOT
in the sample test paper, then checks the parser segments them correctly.

    python -m pytest tests/ -q        (or)        python tests/test_parser.py
"""
import io, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.shared import Inches
from qbg.parser import parse_docx
from qbg.builder import QuestionDocBuilder

SAMPLE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_test_paper.docx")


def make(paras):
    """paras: list of str | ('table', rows) | ('image',) | ('br', [lines])"""
    d = Document()
    for p in paras:
        if isinstance(p, tuple) and p[0] == "table":
            rows = p[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.cell(i, j).text = v
        elif isinstance(p, tuple) and p[0] == "image":
            from PIL import Image
            buf = io.BytesIO(); Image.new("RGB", (40, 20), (200, 30, 30)).save(buf, "PNG"); buf.seek(0)
            d.add_paragraph().add_run().add_picture(buf, width=Inches(1))
        elif isinstance(p, tuple) and p[0] == "br":
            para = d.add_paragraph(); run = para.add_run()
            for k, line in enumerate(p[1]):
                if k: run.add_break()
                run.add_text(line)
        else:
            d.add_paragraph(p)
    out = io.BytesIO(); d.save(out); return out.getvalue()


def opts(q):
    return [" ".join(l.text for l in o) for o in q.options]


def test_sample_paper():
    if not os.path.exists(SAMPLE):
        return
    qs, keys = parse_docx(open(SAMPLE, "rb").read())
    assert len(qs) == 121 and keys == 121
    assert sum(q.ok for q in qs) == 120          # last one has no explanation in source
    assert [q.source_number for q in qs[:120]] == list(range(1, 121))
    assert qs[4].correct_index == 0 and "Necessary that the act abetted" in qs[4].body_text  # type 3
    assert "COMPETITION ACT" not in "\n".join(x.text for x in qs[114].solution)
    b = QuestionDocBuilder("Judiciary")
    for q in qs:
        if q.ok:
            assert b.build(q)[:2] == b"PK"


def test_q_prefix_and_inline_explanation():
    qs, _ = parse_docx(make([
        "Q1. Which court is the apex court?",
        "a) High Court", "b) Supreme Court", "c) District Court", "d) Tribunal",
        "Correct Answer: Option B",
        "Explanation: Article 124 establishes it.",
        "Q2. Capital of India?",
        "(a) Mumbai", "(b) Delhi", "(c) Agra", "(d) Pune",
        "Ans: (b) Explanation: New Delhi is the capital.",
    ]))
    assert len(qs) == 2 and all(q.ok for q in qs)
    assert qs[0].correct_index == 1 and qs[1].correct_index == 1
    assert qs[1].solution[0].text == "New Delhi is the capital."


def test_numbered_list_inside_explanation():
    qs, _ = parse_docx(make([
        "1. First question?", "(a) w", "(b) x", "(c) y", "(d) z", "Ans. (a)",
        "Explanation:", "1. Point one of the explanation.", "2. Point two of the explanation.",
        "2. Second question?", "(a) w", "(b) x", "(c) y", "(d) z", "Ans. (c)",
        "Explanation: fine.",
    ]))
    assert len(qs) == 2, [q.body_text for q in qs]
    assert len(qs[0].solution) == 2 and qs[1].body_text == "Second question?"


def test_explanation_list_then_question_with_statements():
    qs, _ = parse_docx(make([
        "1. Q one?", "(a) w", "(b) x", "(c) y", "(d) z", "Ans. (b)",
        "Explanation:", "1. reason one", "2. reason two",
        "2. Consider the following:", "1. Statement one", "2. Statement two",
        "Which is correct?", "(a) 1 only", "(b) 2 only", "(c) Both", "(d) Neither", "Ans. (c)",
        "Explanation: both are true.",
        "3. Q three?", "(a) w", "(b) x", "(c) y", "(d) z", "Ans. (d)", "Explanation: z.",
    ]))
    assert [q.source_number for q in qs] == [1, 2, 3]
    assert "Statement two" in qs[1].body_text and qs[1].correct_index == 2
    assert [l.text for l in qs[0].solution] == ["1. reason one", "2. reason two"]


def test_type3_statements_directly_above_options():
    qs, _ = parse_docx(make([
        "1. Consider the statements:",
        "(a) Statement A", "(b) Statement B", "(c) Statement C", "(d) Statement D",
        "(a) Only (a)", "(b) Only (b)", "(c) (a) and (c)", "(d) All",
        "Ans. (c)", "Explanation: x",
    ]))
    q = qs[0]
    assert opts(q) == ["Only (a)", "Only (b)", "(a) and (c)", "All"] and q.correct_index == 2
    assert "(d) Statement D" in q.body_text


def test_options_in_table_and_numeric_options():
    qs, _ = parse_docx(make([
        "1. Pick one:", ("table", [["(a) Alpha", "(b) Beta"], ["(c) Gamma", "(d) Delta"]]),
        "Ans. (d)", "Explanation: delta.",
        "2. Numeric style?", "(1) one", "(2) two", "(3) three", "(4) four",
        "Ans. (3)", "Explanation: three.",
    ]))
    assert opts(qs[0]) == ["Alpha", "Beta", "Gamma", "Delta"] and qs[0].correct_index == 3
    assert opts(qs[1]) == ["one", "two", "three", "four"] and qs[1].correct_index == 2


def test_missing_answer_and_doc_starting_mid_paper():
    qs, _ = parse_docx(make([
        "INSTRUCTIONS", "1. Read every question carefully.", "2. No negative marking.",
        "51. First real question?", "(a) w", "(b) x", "(c) y", "(d) z",
        "Explanation: answer key forgotten.",
        "52. Second?", "(a) w", "(b) x", "(c) y", "(d) z", "Answer: b", "Explanation: x.",
    ]))
    assert [q.source_number for q in qs] == [51, 52]
    assert not qs[0].ok and "answer key" in qs[0].remarks()
    assert qs[1].ok and qs[1].correct_index == 1


def test_five_options_heading_image_and_softbreaks():
    qs, _ = parse_docx(make([
        "1. Look at the figure:", ("image",),
        "(a) A", "(b) B", "(c) C", "(d) D", "(e) None",
        "Ans. e", "Explanation: none.", "PART B - EVIDENCE",
        ("br", ["2. Packed in one paragraph?", "A. yes", "B. no", "C. maybe", "D. never",
                "Answer: A", "Explanation: one paragraph."]),
    ]))
    assert len(qs) == 2
    assert len(qs[0].options) == 5 and qs[0].correct_index == 4
    assert any(getattr(l, "has_image", False) for l in qs[0].body)
    assert [l.text for l in qs[0].solution] == ["none."]
    assert qs[1].ok and opts(qs[1]) == ["yes", "no", "maybe", "never"]
    QuestionDocBuilder("Judiciary").build(qs[0])  # image must embed without error


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try:
                fn(); print("PASS", name)
            except AssertionError as e:
                fails += 1; print("FAIL", name, e)
    sys.exit(1 if fails else 0)
