"""End-to-end pipeline test against an in-memory fake of Google Drive."""
import io, json, os, sys, threading, time, random
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
from qbg import pipeline
from qbg.google_services import COUNTER_NAME, XLSX_MIME
from qbg.pipeline import FileTask, Job, to_excel, to_csv

SAMPLE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_test_paper.docx")


class FakeDrive:
    def __init__(self, latency=0.0):
        self.root = "ROOT"
        self.files = {}          # id -> dict(name, parent, mime, data, convert)
        self.folders = {}        # (parent, lname) -> id
        self.lock = threading.Lock(); self.latency = latency
    def check_access(self): return "root"
    def ensure_folder(self, name, parent):
        with self.lock:
            return self.folders.setdefault((parent, name.lower()), f"F{len(self.folders)}")
    def find_file(self, name, parent, folder=False):
        if folder:
            return self.folders.get((parent, name.lower()))
        for fid, f in self.files.items():
            if f["name"] == name and f["parent"] == parent:
                return fid
        return None
    def upload_file(self, data, name, parent, mime, convert_to_gdoc=False):
        time.sleep(self.latency * random.random())
        with self.lock:
            fid = f"id{len(self.files)}"
            self.files[fid] = dict(name=name, parent=parent, mime=mime, data=data, convert=convert_to_gdoc)
        return fid, f"https://drive.google.com/file/d/{fid}/view"
    def reserve_serials(self, output_root, prefix, count, floor=0):
        fid = self.find_file(COUNTER_NAME, output_root)
        last = json.loads(self.files[fid]["data"])["last_serial"] if fid else 0
        last = max(last, floor)
        payload = json.dumps({"last_serial": last + count}).encode()
        if fid: self.files[fid]["data"] = payload
        else: self.upload_file(payload, COUNTER_NAME, output_root, "application/json")
        return last + 1
    def named(self, suffix):
        return {f["name"]: f for f in self.files.values() if f["name"].endswith(suffix)}


def run(job):
    job.start(); job.thread.join(timeout=120); return job


def test_full_run_drive_layout_and_serials():
    pipeline._SERIAL_HIGH_WATER["n"] = 0
    data = open(SAMPLE, "rb").read()
    d = FakeDrive(latency=0.01)
    tasks = [FileTask(0, "TEST_5.docx", data, "Judiciary", "Full Length Test 5"),
             FileTask(1, "bad.docx", b"not a docx", "Judiciary", "Broken file")]
    j = run(Job(tasks, lambda: d, max_workers=12))
    assert j.status == "done", (j.status, j.error, j.log[-5:])
    assert j.total_detected == 121 and j.formatted == 120

    docs = d.named(".docx")
    qdocs = sorted(n for n in docs if n.startswith("QBGDOC"))
    assert qdocs[0] == "QBGDOC000001.docx" and len(qdocs) == 120
    assert not any(f["convert"] for f in docs.values())                 # no Google Doc conversion
    assert "Full Length Test 5.docx" in docs and "Broken file.docx" in docs   # originals in Input

    out = d.folders[("ROOT", "output")]; cat_out = d.folders[(out, "judiciary")]
    q_folder = d.folders[(cat_out, "full length test 5")]
    assert all(docs[n]["parent"] == q_folder for n in qdocs)
    xl = d.named(".xlsx")
    assert set(xl) == {"Full Length Test 5.xlsx", "Broken file.xlsx"}
    assert all(f["parent"] == cat_out and f["mime"] == XLSX_MIME for f in xl.values())

    ws = load_workbook(io.BytesIO(xl["Full Length Test 5.xlsx"]["data"]))["Output"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:7] == ("QBG Category", "File Name", "Question Serial No.", "Original Doc Link",
                           "Question Doc Link", "Status", "Remarks")
    assert len(rows) == 122 and rows[1][2] == "QBGDOC000001" and rows[-1][2] == "QBGDOC000121"
    assert rows[1][4].startswith("https://drive.google.com/file/d/")
    assert sum(r[5] == "Success" for r in rows[1:]) == 120

    # second run (fresh process memory) continues from the Drive counter file
    pipeline._SERIAL_HIGH_WATER["n"] = 0
    run(Job(tasks[:1], lambda: d, max_workers=12))
    assert "QBGDOC000241.docx" in d.named(".docx")      # 242 = the source question with no explanation
    assert "Full Length Test 5 (" in " ".join(d.named(".xlsx"))       # earlier record kept, not overwritten
    assert to_excel(j.results(), {"Run": j.id})[:2] == b"PK" and to_csv(j.results()).startswith(b"\xef\xbb\xbf")
    print(f"  full run: {j.formatted} docs in {j.elapsed:.2f}s")


def test_stop_midway():
    pipeline._SERIAL_HIGH_WATER["n"] = 0
    data = open(SAMPLE, "rb").read()
    d = FakeDrive(latency=0.15)
    j = Job([FileTask(0, "a.docx", data, "Judiciary", "A")], lambda: d, max_workers=4)
    j.start()
    while j.formatted < 8: time.sleep(0.05)
    j.stop(); j.thread.join(timeout=60)
    assert j.status == "stopped" and j.formatted + j.failed == 121
    ws = load_workbook(io.BytesIO(d.named(".xlsx")["A.xlsx"]["data"]))["Output"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    assert len(rows) == 121                                            # every serial accounted for
    stopped = [r for r in rows if r[6] and "stopped by user" in r[6]]
    assert stopped
    print(f"  stop: {j.formatted} formatted before stop, {len(stopped)} marked not processed")


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print("PASS", name)
